from datetime import timedelta
from io import BytesIO
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.generic import TemplateView, View

from apps.core.constants import CASH_FLOW_SECTION_LABELS, NO, YES
from apps.core.mixins import PagePermissionRequiredMixin, PrintContextMixin

from . import statements as fs
from .models import SavedReport
from .services import cash_flow_statement, ledger_integrity

VIEW_TREE = "tree"
VIEW_FLAT = "flat"
VIEW_CHOICES = ((VIEW_TREE, "Chart tree"), (VIEW_FLAT, "Flat by code"))
METHOD_INDIRECT = "indirect"
METHOD_DIRECT = "direct"
METHOD_CHOICES = ((METHOD_INDIRECT, "Indirect"), (METHOD_DIRECT, "Direct"))


def _choice(request, name, choices, default):
    value = request.GET.get(name, default)
    return value if value in dict(choices) else default


class StatementBase(PagePermissionRequiredMixin, PrintContextMixin, TemplateView):
    """Shared by every statement screen: Customise options, drill URLs, header context."""

    template_name = "finance/statement.html"
    title = ""
    url_name = ""
    as_of = False
    default_preset = fs.PRESET_FISCAL_YTD
    toolbar = {"by": True, "compare": True, "customise": True, "fold": True, "pct_income": False}

    def options(self):
        if not hasattr(self, "_options"):
            self._options = fs.resolve_options(self.request, as_of=self.as_of, default_preset=self.default_preset)
        return self._options

    def ledger_url(self, code, column):
        if column is None or not code:
            return ""
        query = {"account_no": code, "date_to": f"{column.end:%Y-%m-%d}", "posted": "1", "back": self.request.get_full_path()}
        if column.start:
            query["date_from"] = f"{column.start:%Y-%m-%d}"
        return f"{reverse('finance:account_ledger')}?{urlencode(query)}"

    def pl_url(self, start, end):
        query = {"preset": fs.PRESET_CUSTOM, "date_from": f"{start:%Y-%m-%d}", "date_to": f"{end:%Y-%m-%d}"}
        return f"{reverse('finance:income_statement')}?{urlencode(query)}"

    def extra_selects(self) -> list:
        return []

    def notice(self) -> dict:
        return {}

    def base_context(self):
        query = self.request.GET.copy()
        query.pop("format", None)
        query.pop("page", None)
        encoded = query.urlencode()
        return {
            "title": self.title,
            "options": self.options(),
            "as_of_report": self.as_of,
            "toolbar": self.toolbar,
            "extra_selects": self.extra_selects(),
            "preset_choices": fs.PRESET_CHOICES,
            "by_choices": fs.BY_CHOICES,
            "compare_choices": fs.COMPARE_CHOICES,
            "show_choices": fs.SHOW_CHOICES,
            "neg_choices": fs.NEG_CHOICES,
            "export_url": f"{reverse(self.url_name)}?{encoded}&format=xlsx" if encoded else f"{reverse(self.url_name)}?format=xlsx",
            "reset_url": reverse(self.url_name),
            "page_query_prefix": f"{encoded}&" if encoded else "",
            "generated_at": timezone.localtime(),
            "report_key": self.url_name,
            "current_query": encoded,
            "saved_reports": [
                {"report": saved, "url": f"{reverse(self.url_name)}?{saved.querystring}", "mine": saved.user_id == self.request.user.pk}
                for saved in SavedReport.objects.filter(report_key=self.url_name).filter(Q(user=self.request.user) | Q(shared=YES)).select_related("user")
            ],
            **self.notice(),
        }

    def get(self, request, *args, **kwargs):
        if request.GET.get("format") == "xlsx":
            return self.xlsx()
        return super().get(request, *args, **kwargs)

    def _workbook(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        org = self._build_print_context(self.request)["org"]
        book = Workbook()
        sheet = book.active
        sheet.title = self.title[:31]
        for text in (str(org) if org else "", self.title, self.options().label, "Accrual basis"):
            sheet.append([text])
        sheet["A1"].font = Font(bold=True, size=13)
        sheet["A2"].font = Font(bold=True)
        sheet.append([])
        return book, sheet

    def _download(self, book, sheet):
        sheet.append([])
        sheet.append([f"Generated {timezone.localtime():%d %b %Y %H:%M} by {self.request.user}"])
        buffer = BytesIO()
        book.save(buffer)
        response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{self.url_name.split(":")[-1].replace("_", "-")}.xlsx"'
        return response

    def xlsx(self):
        raise NotImplementedError


# ---------------------------------------------------------------------------
# Tree statements: P&L, Balance Sheet, Trial Balance, Cash Flow
# ---------------------------------------------------------------------------

class StatementView(StatementBase):
    """``build()`` returns ``{"rows", "columns", "income"?, "computed_drill"?, "blank_zero"?}``."""

    def build(self, options) -> dict:
        raise NotImplementedError

    def statement(self):
        if not hasattr(self, "_statement"):
            options = self.options()
            built = self.build(options)
            columns = built["columns"]
            rows = fs.present(
                built["rows"], columns, options, income=built.get("income"), drill=self.ledger_url,
                computed_drill=built.get("computed_drill"), blank_zero=built.get("blank_zero", False),
            )
            parents = {row["id"]: row["parent"] for row in rows}
            for row in rows:
                chain, cursor = [], row["parent"]
                while cursor and cursor not in chain:
                    chain.append(cursor)
                    cursor = parents.get(cursor)
                row["ancestors"] = chain
            self._statement = {"rows": rows, "columns": columns}
        return self._statement

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        statement = self.statement()
        context.update(self.base_context())
        context.update({
            "table_template": "components/report/statement_table.html",
            "rows": statement["rows"],
            "columns": statement["columns"],
            "landscape": len(statement["columns"]) > 4,
            "fold_spec": "|".join(f'{row["id"]}:{row["depth"]}' for row in statement["rows"] if row["collapsible"]),
        })
        return context

    def xlsx(self):
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.utils import get_column_letter

        statement = self.statement()
        options = self.options()
        book, sheet = self._workbook()
        sheet.append(["", *[column["label"] for column in statement["columns"]]])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        divisor = 1000 if options.thousands else 1
        places = "#,##0" if options.no_paisa else "#,##0.00"
        money_format = f"{places};({places})" if options.neg == fs.NEG_PAREN else f"{places};-{places}"
        for row in statement["rows"]:
            values = []
            for column, cell in zip(statement["columns"], row["cells"]):
                if cell["value"] is None or (row["collapsible"] and row["kind"] in ("section", "group")):
                    values.append(None)
                elif column["kind"] == "pct":
                    values.append(float(cell["value"]) / 100)
                else:
                    values.append(float(cell["value"]) / divisor)
            sheet.append([("    " * row["depth"]) + row["title"], *values])
            line = sheet.max_row
            sheet.row_dimensions[line].outline_level = min(row["depth"], 7)
            for index, column in enumerate(statement["columns"], start=2):
                sheet.cell(line, index).number_format = "0.00%" if column["kind"] == "pct" else money_format
                sheet.cell(line, index).alignment = Alignment(horizontal="right")
            if row["kind"] != "account":
                for cell in sheet[line]:
                    cell.font = Font(bold=True)
            if row["kind"] in ("total", "grand"):
                bottom = Side(style="double") if row["kind"] == "grand" else Side()
                for index in range(2, len(statement["columns"]) + 2):
                    sheet.cell(line, index).border = Border(top=Side(style="thin"), bottom=bottom)
        sheet.column_dimensions["A"].width = 46
        for index in range(len(statement["columns"])):
            sheet.column_dimensions[get_column_letter(index + 2)].width = 18
        return self._download(book, sheet)


class ProfitAndLossView(StatementView):
    page = "finance.income_statement"
    title = "Profit and Loss"
    url_name = "finance:income_statement"
    toolbar = {**StatementBase.toolbar, "pct_income": True}

    def build(self, options):
        base = fs.split_period(options.start, options.end, options.by)
        compared = [fs.compare_column(column, options.compare, span=(options.start, options.end)) for column in base] if options.compare else []
        data = fs.profit_and_loss(base + compared, show=options.show)
        columns = fs.display_columns(base, compared, options, with_income=True)
        return {"rows": data["rows"], "columns": columns, "income": data["totals"]["income"]}


class BalanceSheetView(StatementView):
    page = "finance.balance_sheet"
    title = "Balance Sheet"
    url_name = "finance:balance_sheet"
    as_of = True

    def build(self, options):
        end = options.end
        if options.by == fs.BY_TOTAL:
            base = [fs.Column("c0", f"{end:%d %b %Y}", None, end)]
        else:
            spans = fs.split_period(fs.fiscal_year_start(end), end, options.by)
            base = [fs.Column(column.key, column.label, None, column.end) for column in spans]
        compared = [fs.compare_column(column, options.compare, point=True) for column in base] if options.compare else []
        data = fs.balance_sheet(base + compared, show=options.show)
        years = fs._fiscal_years()

        def computed_drill(row, column):
            if row["id"] == "net_income":
                return self.pl_url(fs.fiscal_year_start(column.end, years), column.end)
            return ""

        return {"rows": data["rows"], "columns": fs.display_columns(base, compared, options), "computed_drill": computed_drill}


class TrialBalanceView(StatementView):
    page = "finance.trial_balance"
    title = "Trial Balance"
    url_name = "finance:trial_balance"
    toolbar = {"by": False, "compare": False, "customise": True, "fold": True, "pct_income": False}

    def view_mode(self):
        return _choice(self.request, "view", VIEW_CHOICES, VIEW_FLAT)

    def extra_selects(self):
        return [{"name": "view", "label": "View", "choices": VIEW_CHOICES, "value": self.view_mode()}]

    def notice(self):
        return {"integrity": ledger_integrity()}

    def build(self, options):
        data = fs.trial_balance(options.start, options.end, tree=self.view_mode() == VIEW_TREE, show=options.show)
        opening = fs.Column("open", "", None, options.start - timedelta(days=1)) if options.start else None
        period = fs.Column("move", "", options.start, options.end)
        closing = fs.Column("close", "", None, options.end)
        columns = []
        for prefix, label, column in (("o", "Opening", opening), ("p", "Period", period), ("c", "Closing", closing)):
            for side, side_label in (("dr", "Dr"), ("cr", "Cr")):
                spec = {"key": f"{prefix}_{side}", "label": f"{label} {side_label}", "kind": "money"}
                if column:
                    spec["column"] = column
                columns.append(spec)
        return {"rows": data["rows"], "columns": columns, "blank_zero": True}


class CashFlowView(StatementView):
    page = "finance.cash_flow"
    title = "Statement of Cash Flows"
    url_name = "finance:cash_flow"

    def method(self):
        return _choice(self.request, "method", METHOD_CHOICES, METHOD_INDIRECT)

    def extra_selects(self):
        return [{"name": "method", "label": "Method", "choices": METHOD_CHOICES, "value": self.method()}]

    def build(self, options):
        if self.method() == METHOD_DIRECT:
            return self._direct(options)
        base = fs.split_period(options.start, options.end, options.by)
        compared = [fs.compare_column(column, options.compare, span=(options.start, options.end)) for column in base] if options.compare else []
        data = fs.cash_flow(base + compared, show=options.show)

        def computed_drill(row, column):
            if row["id"] == "net_income" and column.start:
                return self.pl_url(column.start, column.end)
            return ""

        return {"rows": data["rows"], "columns": fs.display_columns(base, compared, options), "computed_drill": computed_drill}

    def _direct(self, options):
        data = cash_flow_statement(options.start, options.end)
        rows = []
        for section in data["sections"]:
            key = section["key"]
            rows.append(fs._row("section", 0, CASH_FLOW_SECTION_LABELS[key], {"c0": section["subtotal"]}, row_id=key, collapsible=True))
            for line in section["rows"]:
                rows.append(fs._row("account", 1, line["title"], {"c0": line["amount"]}, code=line["code"], parent=key))
            rows.append(fs._row("total", 0, f"Net cash from {CASH_FLOW_SECTION_LABELS[key].lower()}", {"c0": section["subtotal"]}, row_id=f"{key}-total", parent=key))
        rows += [
            fs._row("grand", 0, "Net cash increase for period", {"c0": data["net_movement"]}, row_id="net_change"),
            fs._row("computed", 0, "Cash at beginning of period", {"c0": data["opening"]}, row_id="cash_begin"),
            fs._row("grand", 0, "Cash at end of period", {"c0": data["closing"]}, row_id="cash_end"),
        ]
        if data["difference"]:
            rows.append(fs._row("error", 0, "Difference to cash and bank balance", {"c0": data["difference"]}, row_id="cash_difference"))
        column = fs.Column("c0", "Total", options.start, options.end)
        return {"rows": rows, "columns": [{"key": "c0", "label": "Total", "kind": "money", "column": column}]}


class ChangesInEquityView(StatementView):
    page = "finance.balance_sheet"
    title = "Statement of Changes in Equity"
    url_name = "finance:report_changes_in_equity"

    def build(self, options):
        base = fs.split_period(options.start, options.end, options.by)
        compared = [fs.compare_column(column, options.compare, span=(options.start, options.end)) for column in base] if options.compare else []
        data = fs.changes_in_equity(base + compared, show=options.show)

        def computed_drill(row, column):
            if row["id"] == "net_income" and column.start:
                return self.pl_url(column.start, column.end)
            return ""

        return {"rows": data["rows"], "columns": fs.display_columns(base, compared, options), "computed_drill": computed_drill}


# ---------------------------------------------------------------------------
# Ledger-shaped detail: General Ledger, P&L Detail, Balance Sheet Detail
# ---------------------------------------------------------------------------

class LedgerDetailView(StatementBase):
    toolbar = {"by": False, "compare": False, "customise": False, "fold": False, "pct_income": False}
    accounts_per_page = 25
    amount_mode = False
    exclude_close = False
    with_opening = True
    account_types = None

    def account_filter(self):
        return (self.request.GET.get("account_no") or "").strip()

    def extra_selects(self):
        chart = self.chart()
        choices = [
            (code, f"{code} · {node.title}") for code, node in sorted(chart.by_code.items())
            if not node.children and (self.account_types is None or node.account_type in self.account_types)
        ]
        return [{"name": "account_no", "label": "Account", "choices": [("", "All"), *choices], "value": self.account_filter(), "searchable": True}]

    def chart(self):
        if not hasattr(self, "_chart"):
            self._chart = fs.load_chart()
        return self._chart

    def codes(self):
        chart = self.chart()
        code = self.account_filter()
        if code:
            return {code}
        if self.account_types is None:
            return None
        return chart.codes_of(*self.account_types)

    def accounts(self):
        if not hasattr(self, "_accounts"):
            options = self.options()
            self._accounts = fs.ledger(
                codes=self.codes(), start=options.start, end=options.end, exclude_close=self.exclude_close,
                with_opening=self.with_opening, chart=self.chart(),
            )
        return self._accounts

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(self.base_context())
        accounts = self.accounts()
        paginator = Paginator(accounts, self.accounts_per_page)
        page_obj = paginator.get_page(self.request.GET.get("page"))
        options = self.options()
        column = fs.Column("p", "", options.start, options.end)
        for account in page_obj.object_list:
            account["url"] = self.ledger_url(account["code"], column)
        context.update({
            "table_template": "components/report/ledger_table.html",
            "accounts": page_obj.object_list,
            "page_obj": page_obj,
            "paginator": paginator,
            "is_paginated": paginator.num_pages > 1,
            "amount_mode": self.amount_mode,
            "with_opening": self.with_opening,
            "landscape": True,
            "fold_spec": "",
            "back_query": self.request.get_full_path(),
        })
        return context

    def xlsx(self):
        from openpyxl.styles import Font

        book, sheet = self._workbook()
        money_columns = ["Amount", "Balance"] if self.amount_mode else ["Debit", "Credit", "Balance"]
        sheet.append(["Date", "Type", "No.", "Party", "Memo", "Split", *money_columns])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        fmt = "#,##0.00;(#,##0.00)"
        for account in self.accounts():
            sheet.append([f'{account["code"]} · {account["title"]}'])
            sheet[sheet.max_row][0].font = Font(bold=True)
            if self.with_opening:
                sheet.append(["", "", "", "", "Opening balance", "", *([None] * (len(money_columns) - 1)), float(account["opening"])])
                sheet.row_dimensions[sheet.max_row].outline_level = 1
            for row in account["rows"]:
                figures = [row["amount"], row["balance"]] if self.amount_mode else [row["debit"] or None, row["credit"] or None, row["balance"]]
                sheet.append([row["date"], row["type"], row["number"], row["party"], row["memo"], row["split"], *[float(v) if v is not None else None for v in figures]])
                line = sheet.max_row
                sheet.row_dimensions[line].outline_level = 1
                sheet.cell(line, 1).number_format = "dd-mm-yyyy"
                for index in range(7, 7 + len(money_columns)):
                    sheet.cell(line, index).number_format = fmt
            totals = [account["movement"], account["closing"]] if self.amount_mode else [account["debit"], account["credit"], account["closing"]]
            sheet.append(["", "", "", "", f'Total {account["title"]}', "", *[float(v) for v in totals]])
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
                cell.number_format = fmt
        for letter, width in zip("ABCDEFGHI", (12, 14, 14, 24, 36, 24, 16, 16, 16)):
            sheet.column_dimensions[letter].width = width
        return self._download(book, sheet)


class GeneralLedgerView(LedgerDetailView):
    page = "reports.accounts"
    title = "General Ledger"
    url_name = "finance:report_general_ledger"


class ProfitAndLossDetailView(LedgerDetailView):
    page = "finance.income_statement"
    title = "Profit and Loss Detail"
    url_name = "finance:report_pl_detail"
    amount_mode = True
    exclude_close = True
    with_opening = False
    account_types = fs.PL_TYPES


class BalanceSheetDetailView(LedgerDetailView):
    page = "finance.balance_sheet"
    title = "Balance Sheet Detail"
    url_name = "finance:report_bs_detail"
    amount_mode = True
    account_types = (fs.ACCOUNT_TYPE_ASSET, fs.ACCOUNT_TYPE_LIABILITY, fs.ACCOUNT_TYPE_CAPITAL)


class SavedReportCreateView(LoginRequiredMixin, View):
    def post(self, request):
        report_key = request.POST.get("report_key", "")
        name = (request.POST.get("name") or "").strip()[:120]
        back = reverse(report_key) if report_key in SAVABLE_REPORTS else reverse("finance:income_statement")
        query = request.POST.get("querystring", "")
        if report_key not in SAVABLE_REPORTS or not name:
            return redirect(f"{back}?{query}")
        SavedReport.objects.create(
            user=request.user, report_key=report_key, name=name, querystring=query,
            shared=YES if request.POST.get("shared") == "1" else NO, created_by=request.user, updated_by=request.user,
        )
        messages.success(request, f"Saved \"{name}\".")
        return redirect(f"{back}?{query}")


class SavedReportDeleteView(LoginRequiredMixin, View):
    def post(self, request, pk):
        saved = get_object_or_404(SavedReport, pk=pk, user=request.user)
        saved.soft_delete(user=request.user)
        messages.success(request, f"Removed \"{saved.name}\".")
        return redirect(f"{reverse(saved.report_key)}?{request.POST.get('querystring', '')}")


SAVABLE_REPORTS = {
    view.url_name for view in (
        ProfitAndLossView, BalanceSheetView, TrialBalanceView, CashFlowView, ChangesInEquityView,
        GeneralLedgerView, ProfitAndLossDetailView, BalanceSheetDetailView,
    )
}
